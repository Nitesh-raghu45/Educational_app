"""Document ingestion: parse -> chunk -> embed (MiniLM, local) -> upsert to Pinecone -> mark ready."""

import asyncio
import logging
import re
import uuid
from dataclasses import dataclass
from pathlib import Path

from sqlalchemy import delete, select

from ..config import get_settings
from ..db import SessionLocal
from ..models import Document, DocumentChunk, utcnow
from . import providers

logger = logging.getLogger(__name__)


class IngestionError(Exception):
    """A problem with the file itself; the message is shown to the student."""


@dataclass
class Chunk:
    page: int | None
    text: str


def extract_text(path: Path) -> list[tuple[int | None, str]]:
    """Returns (page number or None, text) pairs."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return [(i, page.extract_text() or "") for i, page in enumerate(reader.pages, 1)]
    if ext == ".docx":
        import docx

        doc = docx.Document(str(path))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return [(None, "\n".join(parts))]
    if ext in (".txt", ".md"):
        return [(None, path.read_text(encoding="utf-8", errors="replace"))]
    if ext in (".png", ".jpg", ".jpeg"):
        try:
            import pytesseract
            from PIL import Image

            return [(None, pytesseract.image_to_string(Image.open(path)))]
        except pytesseract.TesseractNotFoundError as exc:
            raise IngestionError("Image text recognition (Tesseract OCR) is not installed on the server.") from exc
    raise IngestionError(f"Unsupported file type: {ext}")


def split_text(text: str, size: int, overlap: int) -> list[str]:
    text = re.sub(r"[ \t\f\v]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text).strip()
    chunks: list[str] = []
    start = 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):
            window = text[start:end]
            # Prefer breaking at a paragraph, then sentence, then line, then word boundary.
            for separator in ("\n\n", ". ", "\n", " "):
                cut = window.rfind(separator)
                if cut > size // 2:
                    end = start + cut + len(separator)
                    break
        piece = text[start:end].strip()
        if piece:
            chunks.append(piece)
        if end >= len(text):
            break
        next_start = max(end - overlap, start + 1)
        space = text.find(" ", next_start, end)
        start = space + 1 if space != -1 else next_start
    return chunks


def chunk_pages(pages: list[tuple[int | None, str]], size: int, overlap: int) -> list[Chunk]:
    return [Chunk(page, piece) for page, text in pages for piece in split_text(text, size, overlap)]


async def ingest_document(document_id: uuid.UUID) -> None:
    settings = get_settings()
    async with SessionLocal() as session:
        doc = await session.get(Document, document_id)
        if doc is None:
            return
        doc.status, doc.error, doc.updated_at = "processing", None, utcnow()
        await session.commit()
        namespace = providers.namespace_for(doc.user_id)
        try:
            pages = await asyncio.to_thread(extract_text, Path(doc.storage_path))
            chunks = chunk_pages(pages, settings.chunk_size, settings.chunk_overlap)
            if not chunks:
                raise IngestionError(
                    "No text could be extracted. Scanned PDFs aren't supported yet; upload page images instead."
                )

            # Re-ingestion: drop previous vectors and rows first.
            old_ids = list(await session.scalars(select(DocumentChunk.id).where(DocumentChunk.document_id == doc.id)))
            if old_ids:
                await asyncio.to_thread(providers.vector_store().delete, namespace, old_ids)
                await session.execute(delete(DocumentChunk).where(DocumentChunk.document_id == doc.id))

            vectors = await asyncio.to_thread(providers.embedder().embed_documents, [c.text for c in chunks])
            rows = [
                DocumentChunk(id=f"{doc.id.hex}-{i}", document_id=doc.id, chunk_index=i, page=c.page, text=c.text)
                for i, c in enumerate(chunks)
            ]
            items = [
                (
                    row.id,
                    vector,
                    {"document_id": str(doc.id), "chunk_index": row.chunk_index, "page": row.page or 0},
                )
                for row, vector in zip(rows, vectors)
            ]
            await asyncio.to_thread(providers.vector_store().upsert, namespace, items)
            session.add_all(rows)
            doc.status, doc.chunk_count, doc.updated_at = "ready", len(rows), utcnow()
            await session.commit()
            logger.info("Ingested document %s (%d chunks)", doc.id, len(rows))
        except Exception as exc:
            logger.exception("Ingestion failed for document %s", document_id)
            await session.rollback()
            doc = await session.get(Document, document_id, populate_existing=True)
            if doc is not None:
                message = str(exc) if isinstance(exc, IngestionError) else "Processing failed. Try uploading again."
                doc.status, doc.error, doc.updated_at = "failed", message, utcnow()
                await session.commit()
